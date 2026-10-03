from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from .. import models
from ..db import get_db
from ..logic import IST, compute_status, today_ist
from ..services import as_utc, current_subscription, get_settings, members_query

router = APIRouter(tags=["backup"])


def _local(dt):
    dt = as_utc(dt)
    return dt.astimezone(IST).replace(tzinfo=None) if dt else None


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for col, header in zip(ws.columns, headers):
        width = max([len(str(header))] + [len(str(c.value)) for c in col if c.value is not None])
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)
    ws.freeze_panes = "A2"


def _date_format(wb: Workbook):
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.is_date:
                    cell.number_format = "DD MMM YYYY"


@router.get("/backup")
def download_backup(db: Session = Depends(get_db)):
    settings = get_settings(db)
    today = today_ist()
    members = db.scalars(members_query().options(selectinload(models.Member.reminders)).order_by(models.Member.id)).all()
    plans = db.scalars(select(models.Plan).order_by(models.Plan.id)).all()
    names = {m.id: m.name for m in members}

    wb = Workbook()
    wb.remove(wb.active)

    member_rows = []
    for m in members:
        sub = current_subscription(m)
        status = compute_status(sub.end_date, today, settings.expiring_window_days).status if sub else ""
        member_rows.append([
            m.id, m.name, m.phone, m.room_or_address, m.notes,
            sub.plan.name if sub else None, sub.end_date if sub else None, status, _local(m.created_at),
        ])
    _sheet(wb, "Members",
           ["ID", "Name", "Phone", "Room / Address", "Notes", "Current plan", "End date", "Status", "Added on"],
           member_rows)

    sub_rows = []
    for m in members:
        for s in sorted(m.subscriptions, key=lambda s: s.start_date):
            sub_rows.append([
                s.id, m.id, m.name, s.plan.name, s.start_date, s.end_date,
                float(s.amount_paid), s.payment_mode.upper(), s.paid_on,
            ])
    _sheet(wb, "Payments",
           ["ID", "Member ID", "Member", "Plan", "Start date", "End date", "Amount (₹)", "Mode", "Paid on"],
           sub_rows)

    _sheet(wb, "Plans", ["ID", "Name", "Duration", "Unit", "Price (₹)", "Active"],
           [[p.id, p.name, p.duration_value, p.duration_unit, float(p.price), "Yes" if p.is_active else "No"]
            for p in plans])

    reminder_rows = [
        [r.id, r.member_id, names.get(r.member_id), r.type, _local(r.sent_at)]
        for m in members for r in m.reminders
    ]
    _sheet(wb, "Reminders", ["ID", "Member ID", "Member", "Type", "Sent at"], sorted(reminder_rows))

    _sheet(wb, "Settings", ["Setting", "Value"], [
        ["Mess name", settings.mess_name],
        ["Expiring window (days)", settings.expiring_window_days],
        ["Expiring message", settings.expiring_template],
        ["Expired message", settings.expired_template],
    ])
    _date_format(wb)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"mess-backup-{today.isoformat()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
